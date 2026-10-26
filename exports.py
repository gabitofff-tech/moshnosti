from copy import copy
from io import BytesIO
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from engine import ROOT, CATALOG

def safe(v):
    return "'"+v if isinstance(v,str) and v[:1] in '=+-@' else v

def export(e,r,number):
    w=load_workbook(ROOT/'templates'/f'table{number}.xlsx');s=w.active
    if number==1:
        # Keep the entire reference classification and its formatting on a separate sheet.
        ref=w.copy_worksheet(s);ref.title='Нормативы потерь'
        for row in CATALOG:
            for shift,col in enumerate([3,6,9],1): ref.cell(row['row'],col,r['calendar']['hours']*shift)
        header,start,end,style,footer=4,6,81,6,[]
        rows=[]
        for i,q in enumerate(r['equipment'],1):
            a=[i,q['name']]
            for k in range(3): a.extend([r['calendar']['hours']*(k+1),q['losses'][k],q['shifts'][k]['fund']])
            rows.append(a)
        s.cell(1,1,f'Действительный фонд времени работы оборудования — {r["period"]}')
        s.cell(5,1,e['name'])
    elif number==2:
        header,start,end,style=10,11,26,11
        footer=['* — расчёт по выбранному оборудованию; нормы в станко-часах.','Разработал:',e['settings']['developer']]
        products=r['products']; count=max(1,len(products));total_col=4+6*count
        # Repeat exactly the six-column product block with copied template styles.
        total_styles=[copy(s.cell(rr,10)._style) for rr in range(8,11)]
        note_styles=[copy(s.cell(rr,11)._style) for rr in range(8,11)]
        orig=[[copy(s.cell(row,col)) for col in range(3,9)] for row in range(8,11)]
        for merged in list(s.merged_cells):
            if merged.min_col>=3 and merged.min_row>=8: s.unmerge_cells(str(merged))
        for pidx,p in enumerate(products):
            base=3+6*pidx
            for rr in range(8,11):
                for c in range(6):
                    cell=s.cell(rr,base+c);src=orig[rr-8][c]
                    cell._style=copy(src._style);cell.value=src.value
                    s.column_dimensions[get_column_letter(base+c)].width=s.column_dimensions[get_column_letter(3+c)].width
            s.merge_cells(start_row=8,start_column=base,end_row=8,end_column=base+5);s.cell(8,base,safe(p['name']))
        s.merge_cells(start_row=8,start_column=total_col-1,end_row=9,end_column=total_col-1);s.cell(8,total_col-1,'И т. д.')
        for rr in range(8,11):
            s.cell(rr,total_col)._style=copy(total_styles[rr-8]);s.cell(rr,total_col+1)._style=copy(note_styles[rr-8])
        s.column_dimensions[get_column_letter(total_col)].width=s.column_dimensions['J'].width
        s.column_dimensions[get_column_letter(total_col+1)].width=s.column_dimensions['K'].width
        for col,label in [(total_col,'Суммарная трудоёмкость, ст-час'),(total_col+1,'Примечание')]:
            s.merge_cells(start_row=8,start_column=col,end_row=9,end_column=col);s.cell(8,col,label)
        for col in range(1,total_col+2): s.cell(10,col,col)
        rows=[]
        for i,q in enumerate(r['equipment'],1):
            by=[[d for d in r['details'] if d['equipment']==q['id'] and d['product']==p['id']] for p in products]
            for j in range(max([len(x) for x in by]+[1])):
                a=[i if j==0 else None,q['name'] if j==0 else None]
                for ds in by:
                    d=ds[j] if j<len(ds) else None
                    a.extend([d['node'],d['qty'],None if d['minutes'] is None else d['minutes']/60,d['unit'],d['plan'],d['demand']] if d else [None]*6)
                a.extend([None,sum(ds[j]['demand'] for ds in by if j<len(ds) and ds[j]['demand'] is not None),None]); rows.append(a)
            a=[None,'Всего:']
            for ds in by: a.extend([None]*5+[sum(d['demand'] for d in ds if d['demand'] is not None)])
            a.extend([None,q['demand'],None]);rows.append(a)
        for m in list(s.merged_cells):
            if m.min_row==7:s.unmerge_cells(str(m))
        s.merge_cells(start_row=7,start_column=1,end_row=7,end_column=total_col+1)
        s.print_area=f'A1:{get_column_letter(total_col+1)}{start+len(rows)+6}'
    elif number==3:
        header,start,end,style=11,12,24,12
        footer=[f'Максимальный принятый коэффициент загрузки: {e["settings"]["target"]}','* — принятое количество: округление вверх при дробной части > '+str(e['settings']['rounding']),'Разработал:',e['settings']['developer']]
        rows=[]
        for i,q in enumerate(r['equipment'],1):
            a=[i,q['name'],q['count']]
            for sh in q['shifts']: a.extend([sh['fund'],sh['total']])
            a.append(q['demand']);a.extend(x['load'] for x in q['shifts'])
            for sh in q['shifts']:a.extend([sh['need'],sh['accepted']])
            a.extend(x['balance'] for x in q['shifts']);a.append(None);rows.append(a)
    else:
        header,start,end,style=10,11,15,11
        footer=['Постоянная структура программы: общий парк не выделяется каждому изделию заново.','Разработал:',e['settings']['developer']]
        rows=[]
        for i,p in enumerate(r['products'],1):
            a=[i,p['name'],p['plan']]
            for sh in p['shifts']: a.extend([sh['load'],sh['capacity']])
            a.append(None);rows.append(a)
    styles=[copy(s.cell(style,c)._style) for c in range(1,max(s.max_column,max([len(a) for a in rows]+[1]))+1)]
    footer_style=copy(s.cell(25 if number==2 else 23 if number==3 else 13,2)._style) if number!=1 else None
    height=s.row_dimensions[style].height
    for merged in list(s.merged_cells):
        if merged.min_row>=start:s.unmerge_cells(str(merged))
    for row in s.iter_rows(min_row=start):
        for cell in row:
            if not isinstance(cell,MergedCell): cell.value=None;cell._style=None
    for i,values in enumerate(rows,start):
        s.row_dimensions[i].height=height
        for c,v in enumerate(values,1):
            cell=s.cell(i,c,safe(v));cell._style=copy(styles[c-1])
            if isinstance(v,float):cell.number_format='0.00'
    for i,text in enumerate(footer,start+len(rows)+2):
        cell=s.cell(i,2,safe(text));cell._style=copy(footer_style)
    if number!=1:
        s.cell(7,1,f'{s.cell(7,1).value.split(" на ")[0]} на {r["period"]} — {e["name"]}')
        s.cell(2,2,safe(e['settings']['approver_role']));s.cell(4,2,safe('______________________ '+e['settings']['approver']))
    if '-' in r['period']:
        for row in s.iter_rows(max_row=header):
            for cell in row:
                if isinstance(cell.value,str):cell.value=cell.value.replace('годовой','месячный').replace('годового','месячного')
    last=start+len(rows)+len(footer)+4
    s.print_area=f'A1:{get_column_letter(max(len(a) for a in rows) if rows else 11)}{last}'
    s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    s.print_title_rows=f'1:{header}'
    meta=w.create_sheet('Паспорт расчёта')
    import datetime
    for row in [('Предприятие',e['name']),('Период',r['period']),('Методика','1.0 — шаблоны пользователя'),('Дата UTC',datetime.datetime.now(datetime.timezone.utc).isoformat()),('Статус',r['status']),('Календарь',r['calendar']['source']),('Режим мощности','Постоянная структура программы'),('Замечания','; '.join(r['issues']))]: meta.append([safe(x) for x in row])
    out=BytesIO();w.save(out);out.seek(0);return out
