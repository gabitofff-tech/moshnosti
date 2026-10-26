"""Cross-enterprise monitoring and capacity-only transfer scenarios (no baseline writes)."""
import math
from copy import deepcopy
from io import BytesIO
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from engine import calculate, calendar_fund


def group_key(eq):
    return str(eq.get('transfer_group') or '').strip().casefold()


def metrics(demand, fund, target):
    known = demand is not None
    return dict(demand=demand, fund=fund, load=demand/fund if known and fund else None,
                reserve=max(fund-demand,0) if known else None,
                deficit=max(demand-fund,0) if known else None,
                planning_capacity=fund*target,
                free=max(fund*target-demand,0) if known else None,
                excess=max(demand-fund*target,0) if known else None)


def summarize(rows, complete=True):
    known = complete and all(r['demand'] is not None for r in rows)
    fund=sum(r['fund'] for r in rows)
    demand=sum(r['demand'] or 0 for r in rows) if known else None
    return dict(demand=demand,known_demand=sum(r['demand'] or 0 for r in rows),fund=fund,
                load=demand/fund if known and fund else None,
                reserve=sum(r['reserve'] for r in rows) if known else None,
                deficit=sum(r['deficit'] for r in rows) if known else None,
                free=sum(r['free'] for r in rows) if known else None,
                excess=sum(r['excess'] for r in rows) if known else None,
                status='CALCULATED' if known else 'INCOMPLETE')


def monitor(enterprises, year, month, shift):
    if not 2000<=year<=2100 or not 0<=month<=12 or shift not in (1,2,3):
        raise ValueError('Неверный период или сменность')
    cal=calendar_fund(year,month)
    rows=[]; plants=[]; operations=[]
    for e in enterprises:
        try:
            result=calculate(e,year,month)
        except (ValueError,KeyError,TypeError) as exc:
            plants.append(dict(id=e['id'],name=e['name'],issues=[str(exc)],rows=[],status='INCOMPLETE'))
            continue
        complete=not result['issues']
        own=[]
        for eq in result['equipment']:
            r=dict(enterprise=e['id'],enterprise_name=e['name'],equipment=eq['id'],name=eq['name'],
                   group=group_key(eq),count=eq['count'],target=e['settings']['target'],
                   **metrics(eq['demand'] if complete else None,eq['shifts'][shift-1]['total'],e['settings']['target']))
            rows.append(r);own.append(r)
        plants.append(dict(id=e['id'],name=e['name'],issues=result['issues'],rows=own,**summarize(own,complete)))
        if complete:
            for d in result['details']:
                q=next(x for x in rows if x['enterprise']==e['id'] and x['equipment']==d['equipment'])
                if d['minutes'] is not None and d['minutes']>0 and d['plan'] is not None:
                    operations.append(dict(enterprise=e['id'],enterprise_name=e['name'],product=d['product'],
                        product_name=d['product_name'],node=d['node_id'],node_name=d['node'],operation_index=d['operation_index'],
                        operation=d['operation'],equipment=d['equipment'],equipment_name=q['name'],group=q['group'],
                        quantity=d['qty']*d['plan'],minutes=d['minutes'],demand=d['demand']))
    groups={}
    for r in rows:
        # Unconfirmed resources never get pooled on a matching display name/category.
        key=('group',r['group']) if r['group'] else ('unconfirmed',r['enterprise'],r['equipment'])
        groups.setdefault(key,[]).append(r)
    grouped=[dict(name=members[0]['group'] or 'Не подтверждена · '+members[0]['name'],confirmed=bool(members[0]['group']),
                  enterprises=len({r['enterprise'] for r in members}),**summarize(members)) for members in groups.values()]
    candidates=[]
    for op in operations:
        src=next(r for r in rows if r['enterprise']==op['enterprise'] and r['equipment']==op['equipment'])
        if not op['group'] or not src['excess'] or op['quantity']<=0:continue
        for dest in rows:
            if dest['enterprise']==op['enterprise'] or dest['group']!=op['group'] or not dest['free']:continue
            qty=min(op['quantity'],src['excess']*60/op['minutes'],dest['free']*60/op['minutes'])
            candidates.append(dict(source=op,target_enterprise=dest['enterprise'],target_name=dest['enterprise_name'],
                                   target_equipment=dest['equipment'],target_equipment_name=dest['name'],
                                   quantity=qty,hours=qty*op['minutes']/60))
    candidates.sort(key=lambda c:c['hours'],reverse=True)
    return dict(period=f'{year}-{month:02d}' if month else str(year),shift=shift,calendar=cal,
                generated_at=datetime.now(timezone.utc).isoformat(),rows=rows,enterprises=plants,groups=grouped,operations=operations,
                candidates=candidates,summary=summarize(rows,bool(plants) and all(p['status']=='CALCULATED' for p in plants)),
                coverage=dict(complete=sum(p['status']=='CALCULATED' for p in plants),total=len(plants)))


def evaluate(base, transfers):
    result=deepcopy(base); rows=result['rows']; moved={};lines=[]
    def finite(x):return isinstance(x,(int,float)) and not isinstance(x,bool) and math.isfinite(x) and x>0
    for t in transfers:
        op=next((o for o in base['operations'] if all(o[k]==t.get(k) for k in ['enterprise','product','node','operation_index'])),None)
        if op is None:raise ValueError('Исходная операция не найдена или данные предприятия неполны. Обновите мониторинг.')
        src=next(r for r in rows if r['enterprise']==op['enterprise'] and r['equipment']==op['equipment'])
        dst=next((r for r in rows if r['enterprise']==t.get('target_enterprise') and r['equipment']==t.get('target_equipment')),None)
        if dst is None or dst['demand'] is None:raise ValueError('Нет подтверждённой загрузки оборудования получателя')
        if src['enterprise']==dst['enterprise']:raise ValueError('Выберите другое предприятие')
        if not src['group'] or src['group']!=dst['group']:raise ValueError('Не подтверждена совместимая группа оборудования')
        qty=t.get('quantity');minutes=t.get('recipient_minutes')
        if not finite(qty) or not finite(minutes):raise ValueError('Количество и норма получателя должны быть положительными числами')
        key=(op['enterprise'],op['product'],op['node'],op['operation_index'])
        moved[key]=moved.get(key,0)+qty
        if moved[key]>op['quantity']+1e-8:raise ValueError('Сумма переносов превышает программу выбранной операции')
        removed=qty*op['minutes']/60;added=qty*minutes/60
        if dst['demand']+added>dst['planning_capacity']+1e-8:
            raise ValueError('Недостаточно резерва у получателя с учётом остальных строк сценария и допустимого коэффициента')
        src.update(metrics(max(0,src['demand']-removed),src['fund'],src['target']))
        dst.update(metrics(dst['demand']+added,dst['fund'],dst['target']))
        lines.append(dict(source=op,target_enterprise=dst['enterprise'],target_name=dst['enterprise_name'],
                          target_equipment=dst['equipment'],quantity=qty,recipient_minutes=minutes,removed=removed,added=added))
    for plant in result['enterprises']:
        own=[r for r in rows if r['enterprise']==plant['id']]
        plant['rows']=own;plant.update(summarize(own,not plant['issues']))
    grouped={}
    for r in rows:
        key=('group',r['group']) if r['group'] else ('unconfirmed',r['enterprise'],r['equipment'])
        grouped.setdefault(key,[]).append(r)
    result['groups']=[dict(name=members[0]['group'] or 'Не подтверждена · '+members[0]['name'],confirmed=bool(members[0]['group']),enterprises=len({r['enterprise'] for r in members}),**summarize(members)) for members in grouped.values()]
    result['summary']=summarize(rows,bool(result['enterprises']) and all(not p['issues'] for p in result['enterprises']))
    return dict(period=base['period'],shift=base['shift'],lines=lines,before=base,after=result,
                note='Оценка только по станко-времени. Оснастка, персонал, сроки, качество и логистика требуют подтверждения. Исходные программы не изменены.')


def workbook(base,scenario=None):
    from exports import safe
    w=Workbook();w.remove(w.active)
    def sheet(name,heads,records):
        s=w.create_sheet(name);s.append([safe(x) for x in heads])
        for record in records:s.append([safe(v) for v in record])
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]:c.fill=PatternFill('solid',fgColor='163E4B');c.font=Font(color='FFFFFF',bold=True)
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width=min(55,max(17,max(len(str(c.value or '')) for c in col)+2))
            for c in col:
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,float):c.number_format='0.00'
        return s
    heads=['Предприятие','Оборудование','Группа совместимости','Количество','Потребность, ст-ч','Фонд, ст-ч','Кз','Допустимый Кз','Резерв до норматива, ст-ч','Дефицит до норматива, ст-ч','Резерв к фонду, ст-ч','Дефицит к фонду, ст-ч']
    def records(data):
        return [[r['enterprise_name'],r['name'],r['group'],r['count'],r['demand'],r['fund'],r['load'],r['target'],r['free'],r['excess'],r['reserve'],r['deficit']] for r in data['rows']]
    sheet('Загрузка предприятий',heads,records(base))
    sheet('Свод по предприятиям',['Предприятие','Статус','Потребность','Фонд','Кз','Резерв до норматива','Дефицит до норматива','Замечания'],[[p['name'],p['status'],p.get('demand'),p.get('fund'),p.get('load'),p.get('free'),p.get('excess'),'; '.join(p['issues'])] for p in base['enterprises']])
    sheet('Совместимые группы',['Группа','Подтверждена','Предприятий','Потребность','Фонд','Кз','Локальные резервы','Локальные дефициты'],[[g['name'],g['confirmed'],g['enterprises'],g['demand'],g['fund'],g['load'],g['free'],g['excess']] for g in base['groups']])
    if scenario:
        sheet('После переноса',heads,records(scenario['after']))
        sheet('Переносы',['Откуда','Изделие','ДСЕ','Операция','Куда','Количество ДСЕ','Норма получателя, мин','Снято, ст-ч','Добавлено, ст-ч'],[[t['source']['enterprise_name'],t['source']['product_name'],t['source']['node_name'],t['source']['operation'],t['target_name'],t['quantity'],t['recipient_minutes'],t['removed'],t['added']] for t in scenario['lines']])
    sheet('Паспорт',['Параметр','Значение'],[['Период',base['period']],['Сменность',base['shift']],['Дата UTC',base['generated_at']],['Полнота',f"{base['coverage']['complete']}/{base['coverage']['total']}"],['Источник календаря',base['calendar']['source']],['Методика','Мониторинг 1.0; Кз = сумма потребностей / сумма фондов; резервы и дефициты суммируются локально без взаимозачёта'],['Переносы','Только расчётный сценарий, не распоряжение. Строки делят один резерв получателя. Кандидаты — независимые варианты, не суммируются.'],['Ограничения','Оснастка, люди, сроки, качество, логистика не проверены; совместимость групп подтверждает пользователь']])
    out=BytesIO();w.save(out);out.seek(0);return out
