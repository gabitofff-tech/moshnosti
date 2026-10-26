"""End-to-end global monitoring, scenario reservations and Excel export."""
import os
import sys
import tempfile
import time
import subprocess
from io import BytesIO
from pathlib import Path
import httpx
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from test_monitoring import plant
root=Path(__file__).resolve().parents[1]
with tempfile.TemporaryDirectory() as tmp:
    process=subprocess.Popen([sys.executable,'-m','uvicorn','app:app','--host','127.0.0.1','--port','8002'],cwd=root,env=dict(os.environ,CAPACITY_DB=tmp+'/db.sqlite'),stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    try:
        base='http://127.0.0.1:8002'
        for _ in range(50):
            try:
                if httpx.get(base+'/').status_code==200:break
            except httpx.TransportError:pass
            time.sleep(.1)
        for name,d in [('a',180),('b',20),('c',180)]:
            created=httpx.post(base+'/api/enterprises',json={'name':name}).json()
            e=plant(name,d);e.update(id=created['id'],revision=1)
            assert httpx.put(base+'/api/enterprises/'+e['id'],json=e).status_code==200
        with sync_playwright() as p:
            browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
            page=browser.new_page(viewport={'width':1550,'height':1100});errors=[];page.on('pageerror',lambda e:errors.append(str(e)))
            page.goto(base);page.locator('[data-action="tab:monitoring"]').click()
            page.locator('#monitorShift').wait_for();page.locator('#month').select_option('1')
            page.wait_for_function("document.querySelector('#content').textContent.includes('2026-01')")
            assert page.locator('.cards').inner_text().count('3 / 3')==1
            page.locator('[data-action="pickTransfer:0"]').click()
            page.wait_for_function("Number(document.querySelector('#transferQty')?.value)===120")
            page.locator('[data-action="addTransfer"]').click()
            page.get_by_role('heading',name='Результат до / после',exact=True).wait_for()
            page.evaluate('window.scrollTo(0,0)')
            page.screenshot(path=str(root/'docs/monitoring.png'),full_page=True)
            with page.expect_download() as download:
                page.locator('[data-action="exportMonitoring"]').click()
            path=download.value.path();w=load_workbook(BytesIO(Path(path).read_bytes()));assert 'После переноса' in w.sheetnames
            page.locator('[data-action="pickTransfer:0"]').click();page.locator('[data-action="addTransfer"]').click()
            page.wait_for_function("document.querySelector('#notice').textContent.includes('Недостаточно резерва')")
            assert page.get_by_role('heading',name='Результат до / после',exact=True).is_visible()
            page.locator('#monitorShift').select_option('2')
            page.wait_for_function("document.querySelector('#content').textContent.includes('Сценарий переноса')")
            assert page.get_by_role('heading',name='Результат до / после',exact=True).count()==0
            page.locator('[data-action="tab:equipment"]').click()
            assert page.locator('[data-path="equipment.0.transfer_group"]').input_value()=='TOK'
            assert not errors,errors
            browser.close()
        print('Monitoring browser: global totals, candidate, before/after, XLSX, shared reserve rejection, shift reset — PASS')
    finally:
        process.terminate();process.wait(timeout=10)
