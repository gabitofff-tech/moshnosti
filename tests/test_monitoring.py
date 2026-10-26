import copy
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import app
from monitoring import monitor,evaluate,workbook
from test_calculations import fixture

def plant(id,demand,group='TOK',count=1):
    e=fixture();e['id']=id;e['name']='Завод '+id;e['settings']['target']=1
    e['equipment']=e['equipment'][:1];e['equipment'][0].update(count=count,transfer_group=group)
    p=e['products'][0];p['plans']['2026-01']=demand*2
    p['nodes']=[{'id':'part','name':'Деталь','code':'ДЕТ','parent':None,'qty':1,'operations':[{'number':'005','name':'Точение','equipment':'lathe','minutes':30,'alternatives':[]}]}]
    return e

def transfer(eid,qty,target='b',minutes=30):
    return dict(enterprise=eid,product='p',node='part',operation_index=0,target_enterprise=target,target_equipment='lathe',quantity=qty,recipient_minutes=minutes)

class MonitoringTests(unittest.TestCase):
    def test_weighted_total_and_local_deficits_not_netted(self):
        # Same shape as 42,000 / 50,000 = 84%, scaled to a monthly calendar.
        base=monitor([plant(str(i),x) for i,x in enumerate([144,72,96,132,60])],2026,1,1)
        self.assertAlmostEqual(base['summary']['load'],.84)
        self.assertEqual(base['summary']['deficit'],36)
        self.assertEqual(base['summary']['reserve'],132)
        self.assertEqual(len(base['groups']),1)
    def test_weighted_not_average_and_shifts(self):
        b=monitor([plant('a',120),plant('b',0,count=3)],2026,1,1)
        self.assertEqual(b['summary']['load'],.25)
        self.assertEqual(monitor([plant('a',120)],2026,1,2)['summary']['load'],.5)
    def test_no_implicit_compatibility(self):
        base=monitor([plant('a',180,''),plant('b',0,'')],2026,1,1)
        self.assertEqual(base['candidates'],[]);self.assertEqual(len(base['groups']),2)
        with self.assertRaisesRegex(ValueError,'совместимая'):evaluate(base,[transfer('a',10)])
    def test_incomplete_enterprise_never_supplies_reserve(self):
        b=plant('b',0);b['products'][0]['plans'].pop('2026-01')
        base=monitor([plant('a',180),b],2026,1,1)
        self.assertEqual(base['coverage'],dict(complete=1,total=2))
        self.assertIsNone(base['summary']['load']);self.assertIsNone(base['rows'][1]['free']);self.assertEqual(base['candidates'],[])
        with self.assertRaises(ValueError):evaluate(base,[transfer('a',20)])
    def test_scenario_preserves_baseline_and_accounts_for_new_norm(self):
        base=monitor([plant('a',180),plant('b',10)],2026,1,1);snapshot=copy.deepcopy(base)
        sc=evaluate(base,[transfer('a',120,minutes=45)])
        self.assertEqual(base,snapshot)
        self.assertEqual(sc['after']['rows'][0]['demand'],120)
        self.assertEqual(sc['after']['rows'][1]['demand'],100)
        self.assertEqual(sc['after']['summary']['excess'],0)
        self.assertEqual(sc['lines'][0]['removed'],60);self.assertEqual(sc['lines'][0]['added'],90)
    def test_shared_reserve_cannot_be_used_twice(self):
        base=monitor([plant('a',180),plant('c',180),plant('b',20)],2026,1,1)
        evaluate(base,[transfer('a',120)])
        with self.assertRaisesRegex(ValueError,'Недостаточно резерва'):evaluate(base,[transfer('a',120),transfer('c',120)])
    def test_source_quantity_cannot_be_assigned_twice(self):
        base=monitor([plant('a',20),plant('b',0,count=10)],2026,1,1)
        with self.assertRaisesRegex(ValueError,'Сумма переносов'):evaluate(base,[transfer('a',30),transfer('a',30)])
    def test_zero_fund_and_incompatible_group(self):
        base=monitor([plant('a',180),plant('b',0,count=0)],2026,1,1)
        self.assertIsNone(base['rows'][1]['load']);self.assertEqual(base['candidates'],[])
        with self.assertRaises(ValueError):evaluate(base,[transfer('a',1)])
        base=monitor([plant('a',180),plant('b',0,'OTHER')],2026,1,1)
        self.assertEqual(base['candidates'],[])
    def test_target_threshold_and_export(self):
        e=plant('a',120);e['settings']['target']=.95
        b=monitor([e,plant('b',0)],2026,1,1)
        self.assertAlmostEqual(b['rows'][0]['excess'],6)
        sc=evaluate(b,[transfer('a',12)])
        w=load_workbook(workbook(b,sc));self.assertIn('После переноса',w.sheetnames)
        self.assertEqual(w['После переноса'].cell(2,5).value,114)
    def test_api_global_scenario_export_no_writes(self):
        with tempfile.TemporaryDirectory() as tmp,patch.object(app,'DB',Path(tmp)/'test.sqlite'):
            c=TestClient(app.app)
            for name,d in [('a',180),('b',0)]:
                created=c.post('/api/enterprises',json={'name':name}).json()
                e=plant(name,d);e.update(id=created['id'],revision=1)
                self.assertEqual(c.put('/api/enterprises/'+e['id'],json=e).status_code,200)
            original=c.get('/api/enterprises').json();a,b=original
            req=dict(year=2026,month=1,shift=1,transfers=[transfer(a['id'],120,b['id'])])
            self.assertEqual(c.get('/api/monitoring?year=2026&month=1').json()['summary']['load'],.75)
            self.assertEqual(c.post('/api/monitoring/scenario',json=req).status_code,200)
            response=c.post('/api/monitoring/export',json=req)
            self.assertEqual(response.status_code,200);self.assertIn('После переноса',load_workbook(BytesIO(response.content)).sheetnames)
            self.assertEqual(c.get('/api/enterprises').json(),original)
            self.assertEqual(c.get('/api/monitoring?shift=4').status_code,422)
            self.assertEqual(c.get('/api/monitoring?year=2027').status_code,422)

if __name__=='__main__':unittest.main()
