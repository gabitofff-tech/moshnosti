import copy
import tempfile
import unittest
from io import BytesIO
from openpyxl import load_workbook
from fastapi.testclient import TestClient
import app
from engine import calculate,validate,calendar_fund
from exports import export

def fixture():
    return {'id':'e','name':'Завод','equipment':[{'id':'lathe','name':'Токарный','count':1,'category':0,'losses':[0,0,0]},{'id':'spare','name':'Резерв','count':2,'category':0,'losses':[0,0,0]}], 'products':[{'id':'p','name':'Изделие','plans':{'2026':100,'2026-01':100},'nodes':[{'id':'a','code':'СБ','name':'Сборка','qty':2,'parent':None,'operations':[]},{'id':'b','code':'ДЕТ','name':'Деталь','qty':3,'parent':'a','operations':[{'number':'005','name':'Токарная','equipment':'lathe','alternatives':['spare'],'minutes':30}]}]}], 'settings':{'target':.95,'rounding':.2,'approver':'ФИО','approver_role':'Директор','developer':'Технолог'}}

class Tests(unittest.TestCase):
 def test_calendar(self):
    self.assertEqual(calendar_fund(2026)['hours'],1972)
    self.assertEqual(calendar_fund(2026)['days'],247)
    self.assertEqual(sum(calendar_fund(2026,m)['hours'] for m in range(1,13)),1972)
    self.assertEqual(calendar_fund(2026,1)['hours'],120)
    with self.assertRaises(ValueError):calendar_fund(2027)
 def test_bom_minutes_alternatives_and_month(self):
    r=calculate(fixture(),2026,1)
    self.assertEqual(r['details'][0]['qty'],6)
    self.assertEqual(r['equipment'][0]['demand'],300)
    self.assertEqual(r['equipment'][1]['demand'],0)
    self.assertEqual(r['equipment'][0]['shifts'][0]['load'],2.5)
    self.assertEqual(r['products'][0]['shifts'][0]['capacity'],38)
 def test_alternative_explicit(self):
    e=fixture();e['products'][0]['nodes'][1]['operations'][0]['equipment']='spare'
    r=calculate(e,2026,1)
    self.assertEqual(r['equipment'][0]['demand'],0)
    self.assertEqual(r['equipment'][1]['demand'],300)
 def test_cycles_and_unknown(self):
    e=fixture();e['products'][0]['nodes'][0]['parent']='b'
    with self.assertRaisesRegex(ValueError,'Цикл'):validate(e)
    e=fixture();e['products'][0]['nodes'][1]['operations'][0]['minutes']=None
    r=calculate(e,2026,1)
    self.assertEqual(r['status'],'INCOMPLETE');self.assertIsNone(r['equipment'][0]['demand'])
 def test_zero_fund_and_missing_plan(self):
    e=fixture();e['equipment'][0]['count']=0
    r=calculate(e,2026,1)
    self.assertIsNone(r['equipment'][0]['shifts'][0]['load'])
    self.assertIsNone(r['products'][0]['shifts'][0]['capacity'])
    self.assertEqual(calculate(e,2026,2)['status'],'INCOMPLETE')
 def test_assortment_shared_park(self):
    e=fixture();p=copy.deepcopy(e['products'][0]);p['id']='p2';p['name']='Второе';e['products'].append(p)
    r=calculate(e,2026,1)
    self.assertEqual(r['equipment'][0]['demand'],600)
    self.assertEqual([p['shifts'][0]['capacity'] for p in r['products']],[19,19])
 def test_exports_many_products_equipment(self):
    e=fixture()
    for i in range(7):
     q=copy.deepcopy(e['equipment'][0]);q['id']=str(i);e['equipment'].append(q)
    for i in range(3):
     p=copy.deepcopy(e['products'][0]);p['id']=str(i);e['products'].append(p)
    r=calculate(e,2026,1)
    for n in range(1,5):
     w=load_workbook(export(e,r,n));s=w.active
     self.assertIn('Паспорт расчёта',w.sheetnames)
     self.assertIn('2026-01',s.cell(1 if n==1 else 7,1).value)
     self.assertTrue(s.cell(6 if n==1 else 12 if n==3 else 11,2).has_style)
    w=load_workbook(export(e,r,2));self.assertEqual(w.active.cell(8,21).value,e['products'][3]['name'])
 def test_api_save_conflict_export(self):
    with tempfile.TemporaryDirectory() as d:
     original=app.DB;app.DB=__import__('pathlib').Path(d)/'db.sqlite'
     try:
      c=TestClient(app.app)
      e=c.post('/api/enterprises',json={'name':'Тестовый завод'}).json()
      f=fixture();f['id']=e['id'];f['revision']=1
      self.assertEqual(c.put('/api/enterprises/'+e['id'],json=f).status_code,200)
      self.assertEqual(c.put('/api/enterprises/'+e['id'],json=f).status_code,409)
      for n in range(1,5):
       response=c.get(f'/api/export/{e["id"]}/{n}?year=2026&month=1')
       self.assertEqual(response.status_code,200,response.text[:50] if response.status_code!=200 else '')
       self.assertTrue(load_workbook(BytesIO(response.content)).active)
      self.assertEqual(c.get(f'/api/export/{e["id"]}/2?year=2026&month=2').status_code,422)
     finally:app.DB=original

if __name__=='__main__':unittest.main()
